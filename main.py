import os
import sys
import subprocess
import threading
import zipfile
import re
import shutil
import difflib
from tkinter import Canvas, messagebox, filedialog

# Автоматическая установка customtkinter, если его нет
try:
    import customtkinter as ctk
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "customtkinter"])
    import customtkinter as ctk

try:
    import pymupdf as fitz
except ImportError:
    import fitz  # PyMuPDF
import pandas as pd
import openpyxl
import gspread
from oauth2client.service_account import ServiceAccountCredentials
import datetime
from difflib import get_close_matches
import pytesseract
from PIL import Image, ImageEnhance, ImageOps

class GradientProgressBar(ctk.CTkFrame):
    def __init__(self, master, height=15, **kwargs):
        super().__init__(master, height=height, corner_radius=10, fg_color="#121212", border_width=1, border_color="#333333", **kwargs)
        self.height = height
        
        self.canvas = Canvas(self, height=self.height, bg="#121212", highlightthickness=0)
        self.canvas.pack(expand=True, fill="both", padx=2, pady=2)
        
        self.progress = 0.0
        self.canvas.bind("<Configure>", self._on_resize)
        
    def _hex_to_rgb(self, hex_color):
        hex_color = hex_color.lstrip('#')
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        
    def _rgb_to_hex(self, rgb):
        return '#{:02x}{:02x}{:02x}'.format(int(rgb[0]), int(rgb[1]), int(rgb[2]))
        
    def _on_resize(self, event):
        self.update_gradient()
        
    def set(self, value):
        self.progress = max(0.0, min(1.0, value))
        self.after(0, self.update_gradient)
        
    def update_gradient(self):
        self.canvas.delete("gradient")
        width = self.canvas.winfo_width()
        if width <= 1:
            return
            
        draw_width = int(width * self.progress)
        
        rgb1 = self._hex_to_rgb("#FF3333") # Red
        rgb2 = self._hex_to_rgb("#33FF33") # Green
        
        for i in range(draw_width):
            r = rgb1[0] + (rgb2[0] - rgb1[0]) * i / max(1, width - 1)
            g = rgb1[1] + (rgb2[1] - rgb1[1]) * i / max(1, width - 1)
            b = rgb1[2] + (rgb2[2] - rgb1[2]) * i / max(1, width - 1)
            color = self._rgb_to_hex((r, g, b))
            self.canvas.create_line(i, 0, i, self.height, fill=color, tags="gradient")

class PDFSplitterApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        
        # Настройка современного дизайна
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue") # Базовая тема, но цвета мы зададим свои
        
        self.title("Умная разбивка PDF-протоколов")
        self.geometry("900x750")
        
        # Фирменные цвета
        self.bg_root = "#181818"
        self.bg_frame = "#242424"
        self.accent_col = "#FF7B00"
        self.accent_hover = "#FF9522"
        self.text_col = "#EAEAEA"
        self.btn_bg = "#3A3A3A"
        self.btn_hover = "#4A4A4A"
        
        self.configure(fg_color=self.bg_root)
        
        self.file_paths = []
        self.output_dir = None
        self.registry_path = None
        self.is_processing = False
        
        self.setup_ui()
        
    def setup_ui(self):
        # Заголовок
        self.title_lbl = ctk.CTkLabel(
            self, text="Умная разбивка PDF-протоколов", 
            font=ctk.CTkFont(family="Segoe UI", size=24, weight="bold"), 
            text_color=self.accent_col
        )
        self.title_lbl.pack(pady=(30, 20))

        # ФРЕЙМ 1: Выбор файла
        self.frame_input = ctk.CTkFrame(self, fg_color=self.bg_frame, corner_radius=10)
        self.frame_input.pack(fill="x", padx=40, pady=(0, 15))
        
        self.btn_select_file = ctk.CTkButton(
            self.frame_input, text="1. Выбрать PDF файлы", command=self.select_file,
            fg_color=self.btn_bg, hover_color=self.btn_hover,
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"),
            corner_radius=8, width=200, height=40
        )
        self.btn_select_file.pack(side="left", padx=20, pady=20)
        
        self.lbl_file = ctk.CTkLabel(
            self.frame_input, text="Файлы не выбраны", 
            text_color="#999999", font=ctk.CTkFont(family="Segoe UI", size=13),
            wraplength=500, justify="left"
        )
        self.lbl_file.pack(side="left", fill="x", expand=True, padx=(0, 20))

        # ФРЕЙМ 2: Выбор папки
        self.frame_output = ctk.CTkFrame(self, fg_color=self.bg_frame, corner_radius=10)
        self.frame_output.pack(fill="x", padx=40, pady=(0, 15))
        
        self.btn_select_dir = ctk.CTkButton(
            self.frame_output, text="2. Папка сохранения", command=self.select_output_dir,
            fg_color=self.btn_bg, hover_color=self.btn_hover,
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"),
            corner_radius=8, width=200, height=40
        )
        self.btn_select_dir.pack(side="left", padx=20, pady=20)
        
        self.lbl_dir = ctk.CTkLabel(
            self.frame_output, text="Сохранить рядом с исходными файлами (по умолчанию)", 
            text_color="#999999", font=ctk.CTkFont(family="Segoe UI", size=13),
            wraplength=500, justify="left"
        )
        self.lbl_dir.pack(side="left", fill="x", expand=True, padx=(0, 20))

        
        # ФРЕЙМ 3: Выбор реестра (Локальный или Google)
        self.frame_registry = ctk.CTkFrame(self, fg_color=self.bg_frame, corner_radius=10)
        self.frame_registry.pack(fill="x", padx=40, pady=(0, 15))
        
        self.registry_mode = ctk.StringVar(value="local")
        
        self.rb_local = ctk.CTkRadioButton(self.frame_registry, text="Локальный Excel", variable=self.registry_mode, value="local", command=self.toggle_registry_mode)
        self.rb_local.pack(side="left", padx=20, pady=10)
        
        self.rb_google = ctk.CTkRadioButton(self.frame_registry, text="Google Таблица", variable=self.registry_mode, value="google", command=self.toggle_registry_mode)
        self.rb_google.pack(side="left", padx=20, pady=10)
        
        # Контейнер для локального файла
        self.frame_local = ctk.CTkFrame(self.frame_registry, fg_color="transparent")
        self.frame_local.pack(fill="x", padx=20, pady=5)
        
        self.btn_select_registry = ctk.CTkButton(
            self.frame_local, text="Выбрать Excel-файл", command=self.select_registry,
            fg_color=self.btn_bg, hover_color=self.btn_hover,
            font=ctk.CTkFont(family="Segoe UI", size=13), corner_radius=8, width=160, height=35
        )
        self.btn_select_registry.pack(side="left", padx=(0, 10))
        
        self.lbl_registry = ctk.CTkLabel(self.frame_local, text="Файл не выбран", text_color="#999999")
        self.lbl_registry.pack(side="left", fill="x", expand=True)

        # Контейнер для Google Sheets
        self.frame_google = ctk.CTkFrame(self.frame_registry, fg_color="transparent")
        # Изначально скрыт
        
        self.entry_google_url = ctk.CTkEntry(self.frame_google, placeholder_text="Вставьте ссылку на Google Таблицу...", width=300)
        self.entry_google_url.pack(side="left", padx=(0, 10), fill="x", expand=True)
        
        self.btn_creds = ctk.CTkButton(
            self.frame_google, text="Выбрать credentials.json", command=self.select_creds,
            fg_color=self.btn_bg, hover_color=self.btn_hover,
            font=ctk.CTkFont(family="Segoe UI", size=13), corner_radius=8, width=180, height=35
        )
        self.btn_creds.pack(side="left")
        
        self.google_creds_path = None

        # ФРЕЙМ 4: Опции

        self.frame_options = ctk.CTkFrame(self, fg_color="transparent")
        self.frame_options.pack(fill="x", padx=40, pady=(5, 10))
        
        self.chk_force_ocr = ctk.CTkCheckBox(
            self.frame_options, text="Глубокое OCR-распознавание (игнорирует текстовый слой, читает как картинку)", 
            fg_color=self.accent_col, hover_color=self.accent_hover,
            font=ctk.CTkFont(family="Segoe UI", size=13), text_color=self.text_col
        )
        self.chk_force_ocr.pack(side="left", padx=10)

        # КНОПКА ЗАПУСКА
        self.btn_start = ctk.CTkButton(
            self, text="Начать обработку", command=self.start_processing,
            fg_color=self.btn_bg, hover_color=self.btn_hover, # Будет оранжевой при активации
            font=ctk.CTkFont(family="Segoe UI", size=16, weight="bold"),
            corner_radius=8, height=50, state="disabled"
        )
        self.btn_start.pack(fill="x", padx=40, pady=(15, 10))
        
        # ПРОГРЕСС-БАР
        self.progress_bar = GradientProgressBar(self, height=15)
        self.progress_bar.pack(fill="x", padx=40, pady=(0, 20))
        self.progress_bar.set(0)

        # ЛОГИ
        self.log_label = ctk.CTkLabel(
            self, text="Журнал работы:", 
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"), 
            text_color=self.accent_col
        )
        self.log_label.pack(anchor="w", padx=40, pady=(0, 5))
        
        self.log_area = ctk.CTkTextbox(
            self, fg_color="#121212", text_color=self.text_col,
            font=ctk.CTkFont(family="Consolas", size=13), corner_radius=10,
            border_width=1, border_color="#333333"
        )
        self.log_area.pack(fill="both", expand=True, padx=40, pady=(0, 30))
        
    def select_file(self):
        filepaths = filedialog.askopenfilenames(title="Выберите PDF файлы", filetypes=[("PDF Files", "*.pdf")])
        if filepaths:
            self.file_paths = list(filepaths)
            self.lbl_file.configure(text=f"Выбрано файлов: {len(self.file_paths)}", text_color=self.text_col)
            
            self.btn_start.configure(state="normal", fg_color=self.accent_col, hover_color=self.accent_hover)
            
            if not self.output_dir:
                self.output_dir = os.path.dirname(self.file_paths[0])
                self.lbl_dir.configure(text=self.output_dir, text_color=self.text_col)
                
            self.log(f"[*] Выбрано файлов: {len(self.file_paths)}")

    def select_output_dir(self):
        directory = filedialog.askdirectory(title="Выберите папку для сохранения архивов")
        if directory:
            self.output_dir = directory
            self.lbl_dir.configure(text=self.output_dir, text_color=self.text_col)
            self.log(f"[*] Выбрана папка сохранения: {self.output_dir}")
            


    def toggle_registry_mode(self):
        if self.registry_mode.get() == "local":
            self.frame_google.pack_forget()
            self.frame_local.pack(fill="x", padx=20, pady=5)
        else:
            self.frame_local.pack_forget()
            self.frame_google.pack(fill="x", padx=20, pady=5)

    def select_creds(self):
        filepath = filedialog.askopenfilename(title="Выберите credentials.json", filetypes=[("JSON Files", "*.json")])
        if filepath:
            self.google_creds_path = filepath
            self.btn_creds.configure(text="Ключ выбран!", fg_color="#28a745", hover_color="#218838")
            self.log(f"[*] Выбран ключ API: {os.path.basename(filepath)}")

    def select_registry(self):

        filepath = filedialog.askopenfilename(title="Выберите Excel реестр", filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")])
        if filepath:
            self.registry_path = filepath
            self.lbl_registry.configure(text=os.path.basename(filepath), text_color=self.text_col)
            self.log(f"[*] Выбран Excel-реестр: {self.registry_path}")

    def log(self, message):

        self.after(0, self._append_log, message)
        
    def _append_log(self, message):
        self.log_area.insert("end", message + "\n")
        self.log_area.see("end")
        
    def set_buttons_state(self, state):
        self.btn_select_file.configure(state=state)
        self.btn_select_dir.configure(state=state)
        self.btn_select_registry.configure(state=state)
        self.btn_start.configure(state=state)
        self.chk_force_ocr.configure(state=state)
        
        if state == "disabled":
            self.btn_start.configure(fg_color=self.btn_bg)
        else:
            self.btn_start.configure(fg_color=self.accent_col)
            
    def find_tesseract(self):
        try:
            pytesseract.get_tesseract_version()
            return True
        except:
            pass
            
        # Папка, где находится .exe или main.py
        exe_dir = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
        cwd_dir = os.getcwd()
            
        tesseract_paths = [
            os.path.join(exe_dir, "tesseract.exe"),
            os.path.join(exe_dir, "Tesseract-OCR", "tesseract.exe"),
            os.path.join(exe_dir, "tesseract", "tesseract.exe"),
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            os.path.join(os.environ.get("LOCALAPPDATA", ""), r"Programs\Tesseract-OCR\tesseract.exe")
        ]
        
        # Динамический автопоиск tesseract.exe в папке со скриптом и рабочей папке
        import glob
        for search_dir in set([exe_dir, cwd_dir]):
            # Ограничим поиск, чтобы избежать долгого сканирования диска С
            if search_dir != "C:\\" and search_dir != "C:":
                try:
                    found = glob.glob(os.path.join(search_dir, '**', 'tesseract.exe'), recursive=True)
                    tesseract_paths.extend(found)
                except Exception:
                    pass
                    
        for path in tesseract_paths:
            if os.path.exists(path):
                pytesseract.pytesseract.tesseract_cmd = path
                return True
        return False
        
    def clean_protocol_number(self, num_str):
        num_str = num_str.strip()
        
        # 1. Отсекаем все буквенные префиксы (любые варианты СВБ, CBB и прочий мусор)
        core_str = re.sub(r'^[^0-9]*[СC][ВB][БB5]\s*-?', '', num_str, flags=re.IGNORECASE)
        
        # 2. Исправляем буквы, которые OCR принял за цифры (например lO-l707 -> 10-1707)
        replacements = {
            "O": "0", "o": "0", "О": "0", "о": "0",
            "l": "1", "I": "1", "І": "1", "|": "1", "!": "1",
            "З": "3", "б": "6", "В": "8", "B": "8"
        }
        for k, v in replacements.items():
            core_str = core_str.replace(k, v)
            
        # 3. Очищаем весь оставшийся буквенный мусор спереди
        core_str = re.sub(r'^[^0-9]+', '', core_str)
        
        # 4. Восстанавливаем типичные слияния символов в OCR
        # Слэш и двойка часто сливаются в цифру 4 (например /26-1 читается как 46-1)
        core_str = re.sub(r'([0-9]+)46\-1', r'\1_26-1', core_str)
        core_str = re.sub(r'([0-9]+)46\-', r'\1_26-', core_str)
        
        # Специфичное восстановление потерянного префикса, если OCR "съел" начало
        # (например 707_26-1 вместо 10-1707_26-1)
        if core_str.startswith("707") and "26" in core_str:
            core_str = "10-1" + core_str
        
        # 5. Жесткое правило: протоколы начинаются ТОЛЬКО с 10, 12, 28 или 7.
        # Ищем это обязательное начало, за которым идет разделитель (- или _)
        match = re.search(r'(10|12|28|7)[\-_\\]', core_str)
        if match:
            # Отрезаем любой цифровой мусор, который OCR приписал в начале
            core_str = core_str[match.start():]
        else:
            # Если правильного префикса нет, возможно OCR потерял первую цифру
            # Например, 0- вместо 10-, 2- вместо 12-, 8- вместо 28-
            match_broken = re.search(r'(0|2|8)([\-_\\])', core_str)
            if match_broken:
                digit = match_broken.group(1)
                sep = match_broken.group(2)
                prefix = '10' if digit == '0' else ('12' if digit == '2' else '28')
                core_str = prefix + sep + core_str[match_broken.end():]
            else:
                # Если разделителя почему-то нет, ищем только двузначные префиксы. 
                # Одиночную 7 без дефиса не берем, чтобы не обрезать номера вроде 1707 до 707.
                match_any = re.search(r'(10|12|28)', core_str)
                if match_any:
                    core_str = core_str[match_any.start():]
                
        core_str = core_str.replace("/", "_").replace("\\", "_")
                    
        return f"СВБ-{core_str}"

    def start_processing(self):
        if not self.file_paths:
            return
            
        if self.is_processing:
            return
            
        self.use_ocr = self.chk_force_ocr.get()
        
        if not self.use_ocr:
            try:
                # Проверим только первый файл для скорости
                doc = fitz.open(self.file_paths[0])
                has_text = any(doc.load_page(i).get_text("text").strip() for i in range(min(5, len(doc))))
                doc.close()
                
                if not has_text:
                    ans = messagebox.askyesno(
                        "Требуется OCR", 
                        "В первом PDF-файле не найден распознанный текст.\n"
                        "Включить автоматическое распознавание (OCR) для всех файлов?"
                    )
                    if not ans:
                        return
                    self.use_ocr = True
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось прочитать файл: {str(e)}")
                return
                
        if self.use_ocr:
            if not self.find_tesseract():
                messagebox.showinfo(
                    "Tesseract не найден", 
                    "Программа не смогла найти Tesseract OCR.\nУкажите путь к tesseract.exe"
                )
                tess_path = filedialog.askopenfilename(
                    title="Укажите путь к tesseract.exe",
                    filetypes=[("Executable", "tesseract.exe *.exe")]
                )
                if tess_path:
                    pytesseract.pytesseract.tesseract_cmd = tess_path
                else:
                    self.log("[!] Отмена обработки: Tesseract OCR не настроен.")
                    return
        
        self.is_processing = True
        self.set_buttons_state("disabled")
        self.log_area.delete("0.0", "end")
        self.log("=== Начало обработки файлов ===")
        
        thread = threading.Thread(target=self.process_pdfs, daemon=True)
        thread.start()
        
    def process_pdfs(self):
        temp_dir = None
        try:
            project_output_dir = self.output_dir
            temp_dir = os.path.join(project_output_dir, "Сборка_протоколов_temp")
            os.makedirs(temp_dir, exist_ok=True)
            
            all_records = []
            files_to_zip = []
            
            total_files = len(self.file_paths)
            self.progress_bar.set(0)
            
            pattern = re.compile(r"ПРОТОКОЛ\s*(?:№|N|Ne|N2)?\s*([A-Za-zА-Яа-я0-9\-_/\\]+)(?:\s*от\s*([0-9]{2}\.[0-9]{2}\.[0-9]{4}))?", re.IGNORECASE)
            pattern_building = re.compile(r"Корпус\s*([0-9A-Za-zА-Яа-я\-]+)", re.IGNORECASE)
            
            for file_idx, file_path in enumerate(self.file_paths):
                try:
                    self.log(f"\n[{file_idx+1}/{total_files}] Открытие: {os.path.basename(file_path)}")
                    doc = fitz.open(file_path)
                    num_pages = len(doc)
                    
                    protocols = []
                    current_protocol = None
                    
                    for page_num in range(num_pages):
                        # Обновляем прогресс-бар: базовая часть от файлов + часть от текущего файла
                        overall_progress = (file_idx + (page_num / num_pages)) / total_files
                        self.progress_bar.set(overall_progress)
                        
                        page = doc.load_page(page_num)
                        
                        best_match = None
                        p_num, p_date, p_building = None, None, None
                        
                        # ЖЕСТКОЕ ОГРАНИЧЕНИЕ: Максимум 2 страницы на протокол
                        if current_protocol and (page_num - current_protocol["start_page"] == 2):
                            protocols.append(current_protocol)
                            current_protocol = None
                            self.log(f"[*] Стр. {page_num}: принудительное завершение предыдущего протокола (лимит 2 стр.)")
                        
                        if self.use_ocr:
                            pix = page.get_pixmap(dpi=400)
                            base_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                            
                            self.log(f"  Анализ стр. {page_num + 1}/{num_pages}...")
                            
                            for rotation in [0, 180]:
                                if rotation == 180 and best_match and best_match[1]:
                                    break
                                    
                                if rotation == 180:
                                    self.log(f"  [-] Стр. {page_num + 1}: Данные неполные. Переворот на 180 градусов и повтор...")
                                    rotated_img = base_img.rotate(180, expand=True)
                                else:
                                    rotated_img = base_img

                                for attempt in range(5):
                                    if attempt == 0:
                                        img, config = rotated_img, '--psm 3'
                                    elif attempt == 1:
                                        img, config = ImageEnhance.Contrast(ImageOps.grayscale(rotated_img)).enhance(2.0), '--psm 3'
                                    elif attempt == 2:
                                        img, config = rotated_img.convert('L').point(lambda x: 0 if x < 130 else 255, '1'), '--psm 6'
                                    elif attempt == 3:
                                        img, config = rotated_img.convert('L').point(lambda x: 0 if x < 180 else 255, '1'), '--psm 6'
                                    else:
                                        w, h = rotated_img.size
                                        img = rotated_img.crop((0, 0, w, h // 3))
                                        config = '--psm 6'
                                        
                                    text = pytesseract.image_to_string(img, lang="rus+eng", config=config)
                                    header_text = text[:3000].replace('\n', ' ')
                                    
                                    match = pattern.search(header_text)
                                    match_building = pattern_building.search(header_text)
                                    
                                    temp_building = match_building.group(1).strip() if match_building else None
                                    
                                    if match:
                                        prefix = header_text[max(0, match.start() - 30):match.start()].lower()
                                        if "продолжени" in prefix:
                                            continue
                                            
                                        temp_num = self.clean_protocol_number(match.group(1))
                                        temp_date = match.group(2)
                                        
                                        temp_c_date = None
                                        match_c_date = pattern_cdate.search(header_text)
                                        if match_c_date:
                                            temp_c_date = match_c_date.group(1).strip()
                                        
                                        temp_struct = None
                                        match_struct = pattern_struct.search(header_text)
                                        if match_struct:
                                            temp_struct = match_struct.group(1).strip().replace('\n', ' ')
                                        
                                        temp_age = None
                                        match_age = pattern_age.search(header_text)
                                        if match_age:
                                            temp_age = match_age.group(1).strip()
                                        
                                        if temp_date:
                                            final_building = temp_building
                                            if not final_building and best_match and best_match[2]:
                                                final_building = best_match[2]
                                            best_match = (temp_num, temp_date, final_building, temp_struct, temp_c_date, temp_age)
                                            break
                                        elif best_match is None:
                                            best_match = (temp_num, None, temp_building, temp_struct, temp_c_date, temp_age)
                                        else:
                                            if temp_building and not best_match[2]:
                                                best_match = (best_match[0], best_match[1], temp_building, best_match[3], best_match[4], best_match[5])
                                            
                                if best_match and best_match[1]:
                                    break
                                    
                            if best_match:
                                p_num, p_date, p_building, p_structure, p_concreting_date, p_age = best_match
                        else:
                            p_structure, p_concreting_date, p_age = None, None, None
                            text = page.get_text("text").replace('\n', ' ')
                            match = pattern.search(text[:3000])
                            match_building = pattern_building.search(text[:3000])
                            if match_building:
                                p_building = match_building.group(1).strip()
                                
                            match_c_date = pattern_cdate.search(text[:3000])
                            if match_c_date:
                                p_concreting_date = match_c_date.group(1).strip()
                                
                            match_struct = pattern_struct.search(text[:3000])
                            if match_struct:
                                p_structure = match_struct.group(1).strip()
                                
                            match_age = pattern_age.search(text[:3000])
                            if match_age:
                                p_age = match_age.group(1).strip()
                                
                            if match:
                                prefix = text[max(0, match.start() - 30):match.start()].lower()
                                if "продолжени" not in prefix:
                                    p_num = self.clean_protocol_number(match.group(1))
                                    p_date = match.group(2)
                            
                        if p_num:
                            if current_protocol:
                                if current_protocol["number"] == p_num and not current_protocol["uncertain"]:
                                    current_protocol["end_page"] = page_num
                                    if p_building and not current_protocol.get("building"):
                                        current_protocol["building"] = p_building
                                else:
                                    protocols.append(current_protocol)
                                    current_protocol = {
                                        "number": p_num, "date": p_date, "building": p_building,
                                        "structure": p_structure, "concreting_date": p_concreting_date, "age": p_age,
                                        "start_page": page_num, "end_page": page_num,
                                        "uncertain": p_date is None
                                    }
                            else:
                                current_protocol = {
                                    "number": p_num, "date": p_date, "building": p_building,
                                    "structure": p_structure, "concreting_date": p_concreting_date, "age": p_age,
                                    "start_page": page_num, "end_page": page_num,
                                    "uncertain": p_date is None
                                }
                        else:
                            if current_protocol:
                                current_protocol["end_page"] = page_num
                            else:
                                current_protocol = {
                                    "number": "СВБ-НЕ_РАСПОЗНАН", "date": None, "building": None,
                                    "structure": None, "concreting_date": None, "age": None,
                                    "start_page": page_num, "end_page": page_num,
                                    "uncertain": True
                                }
                    
                    if current_protocol:
                        protocols.append(current_protocol)
                        
                    if not protocols:
                        self.log(f"[!] Не найдено ни одного протокола в файле {os.path.basename(file_path)}")
                        doc.close()
                        continue
                        
                    # --- ИДЕАЛЬНОЕ СЛИЯНИЕ ДВУХСТРАНИЧНЫХ ПРОТОКОЛОВ ---
                    merged_protocols = []
                    for prot in protocols:
                        if not merged_protocols:
                            merged_protocols.append(prot)
                            continue
                            
                        prev_prot = merged_protocols[-1]
                        prev_len = prev_prot["end_page"] - prev_prot["start_page"] + 1
                        
                        # Мы можем объединить страницу, только если предыдущая часть - ровно 1 страница
                        can_merge = (prev_len == 1) and (prot["start_page"] == prev_prot["end_page"] + 1)
                        
                        if can_merge:
                            if prot["uncertain"]:
                                prev_prot["end_page"] = prot["end_page"]
                                if prot.get("building") and not prev_prot.get("building"):
                                    prev_prot["building"] = prot["building"]
                                continue
                            
                            if prev_prot["uncertain"] and not prot["uncertain"]:
                                prev_prot["number"] = prot["number"]
                                prev_prot["date"] = prot["date"]
                                prev_prot["end_page"] = prot["end_page"]
                                prev_prot["uncertain"] = False
                                if prot.get("building") and not prev_prot.get("building"):
                                    prev_prot["building"] = prot["building"]
                                continue
                                    
                        merged_protocols.append(prot)
                        
                    protocols = merged_protocols
                    
                    # Формирование файлов для текущего PDF
                    for idx, prot in enumerate(protocols, start=1):
                        clean_num = prot["number"].replace("/", "_").replace("\\", "_")
                        date_str = prot["date"] if prot["date"] else "Без_даты"
                        clean_date = date_str.replace("/", "_").replace("\\", "_")
                        building_str = prot.get("building")
                        building_folder_name = f"Корпус {building_str}" if building_str else "Корпус_Неизвестен"
                        
                        start_p = prot["start_page"]
                        end_p = prot["end_page"]
                        num_p = end_p - start_p + 1
                        
                        page_range_str = f"{start_p+1}" if num_p == 1 else f"{start_p+1}-{end_p+1}"
                        
                        # Определяем папку по префиксу (10, 12, 7_28)
                        prefix_folder = "Неизвестно"
                        match_prefix = re.search(r'СВБ-(10|12|28|7)[\-_\\]', prot["number"])
                        if match_prefix:
                            p_val = match_prefix.group(1)
                            prefix_folder = "7_28" if p_val in ["7", "28"] else p_val
                        else:
                            match_any = re.search(r'СВБ-(10|12|28|7)', prot["number"])
                            if match_any:
                                p_val = match_any.group(1)
                                prefix_folder = "7_28" if p_val in ["7", "28"] else p_val
                        
                        # ФОРМАТ: Протокол_СВБ-10-1807_26-1_от_25.07.2026.pdf
                        # Замечание: убрали _стр_X-Y по просьбе пользователя
                        filename = f"Протокол_{clean_num}_от_{clean_date}.pdf"
                        
                        # Сохраняем в выбранную папку со структурой Корпус -> Префикс
                        target_dir = os.path.join(project_output_dir, building_folder_name, prefix_folder)
                        os.makedirs(target_dir, exist_ok=True)
                        
                        target_filepath = os.path.join(target_dir, filename)
                        
                        # Для предотвращения перезаписи при одинаковых именах из разных файлов
                        counter = 1
                        while os.path.exists(target_filepath):
                            filename_base = f"Протокол_{clean_num}_от_{clean_date}_{counter}.pdf"
                            target_filepath = os.path.join(target_dir, filename_base)
                            filename = filename_base
                            counter += 1
                            
                        temp_filepath = os.path.join(temp_dir, filename)
                        
                        new_doc = fitz.open()
                        new_doc.insert_pdf(doc, from_page=start_p, to_page=end_p)
                        
                        # Сохраняем PDF в иерархию папок
                        new_doc.save(target_filepath)
                        # Сохраняем PDF во временную папку для добавления в ZIP
                        new_doc.save(temp_filepath)
                        
                        new_doc.close()
                        files_to_zip.append(temp_filepath)
                        
                        note = "уверенно" if not prot["uncertain"] else "требует ручной проверки"
                        
                        all_records.append({
                            "Исходный файл": os.path.basename(file_path),
                            "Номер протокола": prot["number"],
                            "Дата протокола": prot["date"] if prot["date"] else "Не найдена",
                            "Корпус": building_str if building_str else "Не найден",
                            "Наименование конструкции": prot.get("structure"),
                            "Дата бетонирования": prot.get("concreting_date"),
                            "Возраст суток": prot.get("age"),
                            "Страницы в исходном PDF": page_range_str,
                            "Количество страниц": num_p,
                            "Имя файла": filename,
                            "Примечание": note,
                            "Префикс_Колонка": prefix_folder
                        })
                        
                    doc.close()
                    
                except Exception as ex:
                    self.log(f"[ОШИБКА] Ошибка при обработке {os.path.basename(file_path)}: {str(ex)}")
            
            # --- Завершение обработки всех файлов ---
            if all_records:
                df = pd.DataFrame(all_records)
                excel_path = os.path.join(temp_dir, "Реестр_разбивки.xlsx")
                df.to_excel(excel_path, index=False)
                files_to_zip.append(excel_path)
                
                # --- ИНТЕГРАЦИЯ В СУЩЕСТВУЮЩИЙ РЕЕСТР ---
                mode = getattr(self, 'registry_mode', None)
                if mode:
                    mode_val = mode.get()
                    unmatched_records = all_records.copy()
                    
                    if mode_val == "local" and getattr(self, 'registry_path', None) and os.path.exists(self.registry_path):
                        self.log("\n[*] Начинаем интеграцию в локальный реестр Excel...")
                        try:
                            wb = openpyxl.load_workbook(self.registry_path)
                            
                            # Перебираем все листы, кроме 'Нераспределенные'
                            sheets_to_process = [sheet for sheet in wb.sheetnames if sheet != "Нераспределенные"]
                            
                            for sheet_name in sheets_to_process:
                                ws = wb[sheet_name]
                                self.log(f"  Проверка листа: {sheet_name}")
                                
                                col_date = 2
                                col_building = 3
                                col_struct = 4
                                col_7_nz = 5
                                col_28_nz = 7
                                col_7_cub = 9
                                col_28_cub = 11
                                
                                for col in range(1, ws.max_column + 1):
                                    val = str(ws.cell(row=1, column=col).value).lower()
                                    if "дата бетонирования" in val: col_date = col
                                    elif "корпус" in val: col_building = col
                                    elif "наименование конструкции" in val: col_struct = col
                                    elif "неразрушающего контроля на 7" in val: col_7_nz = col
                                    elif "неразрушающего контроля на 28" in val: col_28_nz = col
                                    elif "на 7 сутки (кубики)" in val: col_7_cub = col
                                    elif "на 28 сутки (кубики)" in val: col_28_cub = col
                                    
                                still_unmatched = []
                                
                                for rec in unmatched_records:
                                    if rec["Номер протокола"] == "СВБ-НЕ_РАСПОЗНАН":
                                        still_unmatched.append(rec)
                                        continue
                                        
                                    p_num = rec["Номер протокола"]
                                    p_date = rec["Дата протокола"]
                                    p_bldg = str(rec["Корпус"]).strip()
                                    p_strc = str(rec["Наименование конструкции"]).strip() if rec["Наименование конструкции"] else ""
                                    p_cdate = str(rec["Дата бетонирования"]).strip() if rec["Дата бетонирования"] else ""
                                    p_age = str(rec["Возраст суток"]).strip() if rec["Возраст суток"] else ""
                                    p_prefix = rec["Префикс_Колонка"]
                                    
                                    best_row = None
                                    
                                    for r in range(2, ws.max_row + 1):
                                        e_bldg = str(ws.cell(row=r, column=col_building).value or "").strip()
                                        e_cdate = ws.cell(row=r, column=col_date).value
                                        if isinstance(e_cdate, datetime.datetime):
                                            e_cdate = e_cdate.strftime("%d.%m.%Y")
                                        else:
                                            e_cdate = str(e_cdate or "").strip()
                                        
                                        if p_bldg in e_bldg and (not p_cdate or p_cdate == e_cdate or p_cdate.replace('.', '') in e_cdate.replace('.', '')):
                                            e_strc = str(ws.cell(row=r, column=col_struct).value or "").strip()
                                            if p_strc and e_strc:
                                                matches = get_close_matches(p_strc, [e_strc], n=1, cutoff=0.5)
                                                if matches or (p_strc.lower()[:10] in e_strc.lower()):
                                                    best_row = r
                                                    break
                                            else:
                                                best_row = r
                                                break
                                                
                                    if best_row:
                                        target_col = None
                                        if p_prefix == "10":
                                            if p_age == "28": target_col = col_28_nz
                                            else: target_col = col_7_nz
                                        elif p_prefix == "12":
                                            if p_age == "28": target_col = col_28_cub
                                            else: target_col = col_7_cub
                                        elif p_prefix == "7": target_col = col_7_cub
                                        elif p_prefix == "28": target_col = col_28_cub
                                            
                                        if target_col:
                                            entry = f"№{p_num} от {p_date}"
                                            ws.cell(row=best_row, column=target_col).value = entry
                                            
                                            filename = rec["Имя файла"]
                                            rel_path = f"Корпус {p_bldg}/{p_prefix}/{filename}" if p_bldg else f"Корпус_Неизвестен/{p_prefix}/{filename}"
                                            ws.cell(row=best_row, column=target_col).hyperlink = rel_path
                                            ws.cell(row=best_row, column=target_col).style = "Hyperlink"
                                        else:
                                            still_unmatched.append(rec)
                                    else:
                                        still_unmatched.append(rec)
                                
                                unmatched_records = still_unmatched
                                
                            if unmatched_records:
                                if "Нераспределенные" in wb.sheetnames:
                                    ws_un = wb["Нераспределенные"]
                                else:
                                    ws_un = wb.create_sheet("Нераспределенные")
                                    ws_un.append(["Номер протокола", "Дата протокола", "Корпус", "Наименование конструкции", "Дата бетонирования", "Имя файла", "Ожидаемый столбец"])
                                    
                                for rec in unmatched_records:
                                    ws_un.append([
                                        rec["Номер протокола"], rec["Дата протокола"], rec["Корпус"], 
                                        rec.get("Наименование конструкции", ""), rec.get("Дата бетонирования", ""),
                                        rec["Имя файла"], rec["Префикс_Колонка"]
                                    ])
                            
                            out_registry = os.path.join(project_output_dir, "Реестр_обновлен.xlsx")
                            wb.save(out_registry)
                            self.log(f"[✓] Локальный реестр успешно обновлен: {out_registry}")
                        except Exception as e:
                            self.log(f"[ОШИБКА] Не удалось обновить реестр: {str(e)}")
                            
                    elif mode_val == "google" and getattr(self, 'google_creds_path', None) and self.entry_google_url.get():
                        self.log("\n[*] Начинаем интеграцию в Google Таблицу...")
                        try:
                            scope = ['https://spreadsheets.google.com/feeds', 'https://www.googleapis.com/auth/drive']
                            creds = ServiceAccountCredentials.from_json_keyfile_name(self.google_creds_path, scope)
                            client = gspread.authorize(creds)
                            
                            url = self.entry_google_url.get().strip()
                            sh = client.open_by_url(url)
                            
                            for ws in sh.worksheets():
                                if ws.title == "Нераспределенные": continue
                                self.log(f"  Проверка листа: {ws.title}")
                                
                                data = ws.get_all_values()
                                if not data: continue
                                
                                col_date = 2
                                col_building = 3
                                col_struct = 4
                                col_7_nz = 5
                                col_28_nz = 7
                                col_7_cub = 9
                                col_28_cub = 11
                                
                                headers = [h.lower() for h in data[0]]
                                for c_idx, val in enumerate(headers):
                                    if "дата бетонирования" in val: col_date = c_idx + 1
                                    elif "корпус" in val: col_building = c_idx + 1
                                    elif "наименование конструкции" in val: col_struct = c_idx + 1
                                    elif "неразрушающего контроля на 7" in val: col_7_nz = c_idx + 1
                                    elif "неразрушающего контроля на 28" in val: col_28_nz = c_idx + 1
                                    elif "на 7 сутки (кубики)" in val: col_7_cub = c_idx + 1
                                    elif "на 28 сутки (кубики)" in val: col_28_cub = c_idx + 1
                                    
                                still_unmatched = []
                                updates = []
                                
                                for rec in unmatched_records:
                                    if rec["Номер протокола"] == "СВБ-НЕ_РАСПОЗНАН":
                                        still_unmatched.append(rec)
                                        continue
                                        
                                    p_num = rec["Номер протокола"]
                                    p_date = rec["Дата протокола"]
                                    p_bldg = str(rec["Корпус"]).strip()
                                    p_strc = str(rec["Наименование конструкции"]).strip() if rec["Наименование конструкции"] else ""
                                    p_cdate = str(rec["Дата бетонирования"]).strip() if rec["Дата бетонирования"] else ""
                                    p_age = str(rec["Возраст суток"]).strip() if rec["Возраст суток"] else ""
                                    p_prefix = rec["Префикс_Колонка"]
                                    
                                    best_row = None
                                    
                                    for r_idx, row in enumerate(data):
                                        if r_idx == 0: continue
                                        
                                        e_bldg = str(row[col_building-1] if col_building-1 < len(row) else "").strip()
                                        e_cdate = str(row[col_date-1] if col_date-1 < len(row) else "").strip()
                                        
                                        if p_bldg in e_bldg and (not p_cdate or p_cdate == e_cdate or p_cdate.replace('.', '') in e_cdate.replace('.', '')):
                                            e_strc = str(row[col_struct-1] if col_struct-1 < len(row) else "").strip()
                                            if p_strc and e_strc:
                                                matches = get_close_matches(p_strc, [e_strc], n=1, cutoff=0.5)
                                                if matches or (p_strc.lower()[:10] in e_strc.lower()):
                                                    best_row = r_idx + 1
                                                    break
                                            else:
                                                best_row = r_idx + 1
                                                break
                                                
                                    if best_row:
                                        target_col = None
                                        if p_prefix == "10":
                                            if p_age == "28": target_col = col_28_nz
                                            else: target_col = col_7_nz
                                        elif p_prefix == "12":
                                            if p_age == "28": target_col = col_28_cub
                                            else: target_col = col_7_cub
                                        elif p_prefix == "7": target_col = col_7_cub
                                        elif p_prefix == "28": target_col = col_28_cub
                                            
                                        if target_col:
                                            entry = f"№{p_num} от {p_date}"
                                            # Записываем формулу для гиперссылки в Google Sheets
                                            filename = rec["Имя файла"]
                                            rel_path = f"Корпус {p_bldg}/{p_prefix}/{filename}" if p_bldg else f"Корпус_Неизвестен/{p_prefix}/{filename}"
                                            formula = f'=HYPERLINK("{rel_path}"; "{entry}")'
                                            updates.append({'range': f'{gspread.utils.rowcol_to_a1(best_row, target_col)}', 'values': [[formula]]})
                                        else:
                                            still_unmatched.append(rec)
                                    else:
                                        still_unmatched.append(rec)
                                
                                unmatched_records = still_unmatched
                                if updates:
                                    ws.batch_update(updates, value_input_option='USER_ENTERED')
                                
                            if unmatched_records:
                                try:
                                    ws_un = sh.worksheet("Нераспределенные")
                                except gspread.exceptions.WorksheetNotFound:
                                    ws_un = sh.add_worksheet(title="Нераспределенные", rows="100", cols="10")
                                    ws_un.append_row(["Номер протокола", "Дата протокола", "Корпус", "Наименование конструкции", "Дата бетонирования", "Имя файла", "Ожидаемый столбец"])
                                
                                new_rows = []
                                for rec in unmatched_records:
                                    new_rows.append([
                                        rec["Номер протокола"], rec["Дата протокола"], rec["Корпус"], 
                                        rec.get("Наименование конструкции", ""), rec.get("Дата бетонирования", ""),
                                        rec["Имя файла"], rec["Префикс_Колонка"]
                                    ])
                                ws_un.append_rows(new_rows)
                                
                            self.log(f"[✓] Google Таблица успешно обновлена!")
                        except Exception as e:
                            self.log(f"[ОШИБКА] Не удалось обновить Google Таблицу: {str(e)}")
                            
                zip_filename = os.path.join(project_output_dir, "Все_протоколы_архив.zip")
                with zipfile.ZipFile(zip_filename, 'w', zipfile.ZIP_DEFLATED) as zipf:
                    for file in files_to_zip:
                        zipf.write(file, os.path.basename(file))
                        
                # Добавим Excel-файл и в основную папку, чтобы было удобно
                df.to_excel(os.path.join(project_output_dir, "Реестр_разбивки.xlsx"), index=False)
                
                self.progress_bar.set(1.0)
                self.log(f"\n[✓] Готово! \nВсего обработано файлов: {total_files}\nФайлы распакованы в папки по корпусам: {project_output_dir}\nТакже создан ZIP-архив: {zip_filename}")
            else:
                self.log("\n[!] Нет данных для сохранения (ни один протокол не найден).")
                
        except Exception as e:
            self.log(f"\n[КРИТИЧЕСКАЯ ОШИБКА] {str(e)}")
        finally:
            if temp_dir and os.path.exists(temp_dir):
                shutil.rmtree(temp_dir, ignore_errors=True)
            self.is_processing = False
            self.after(0, lambda: self.set_buttons_state("normal"))

if __name__ == "__main__":
    app = PDFSplitterApp()
    app.mainloop()
